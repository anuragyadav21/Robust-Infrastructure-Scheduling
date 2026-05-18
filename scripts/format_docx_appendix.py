#!/usr/bin/env python3
"""Rebuild Appendix A in the thesis DOCX with clean headings, captions, and tables."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "ROBUST OPTIMISATION OF INFRASTRUCTURE SCHEDULING.docx"

HEADER_FILL = "D9E2F3"
BORDER_COLOR = "404040"


def para_text(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


def delete_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER_COLOR)
        borders.append(tag)
    tbl_pr.append(borders)


def shade_row(row, fill: str) -> None:
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)


def write_cell(cell, text: str, *, bold: bool = False, mono: bool = False, size: int = 10) -> None:
    cell.text = text
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if not p.runs:
        run = p.add_run(text)
    else:
        run = p.runs[0]
        for extra in p.runs[1:]:
            extra.text = ""
    run.bold = bold
    run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"
        r_pr = run._element.get_or_add_rPr()
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Consolas")
        r_fonts.set(qn("w:hAnsi"), "Consolas")
        r_pr.insert(0, r_fonts)


def format_data_table(table, headers: list[str], rows: list[list[str]], col_widths: list[float]) -> None:
    try:
        table.style = "Normal Table"
    except KeyError:
        pass
    set_table_borders(table)

    for i, w in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = Inches(w)

    for j, h in enumerate(headers):
        write_cell(table.rows[0].cells[j], h, bold=True, size=10)
        table.rows[0].cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade_row(table.rows[0], HEADER_FILL)

    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        for ci, val in enumerate(row_data):
            mono = ci == 0 and len(row_data) > 1
            write_cell(row.cells[ci], val, mono=mono, size=9)
            if ci == 0:
                row.cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                row.cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    table.autofit = False


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def add_body(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if italic:
        for run in p.runs:
            run.italic = True


def add_section_heading(doc: Document, text: str) -> None:
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)


def add_table_block(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float],
) -> None:
    add_caption(doc, caption)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    format_data_table(table, headers, rows, col_widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def remove_appendix_a(doc: Document) -> None:
    """Remove existing Appendix A section (paragraphs + tables) from document end."""
    body = doc.element.body
    removing = False
    to_remove = []
    for child in list(body):
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        if tag == "p":
            text = para_text(child).strip()
            if text == "Appendix A: File inventory":
                removing = True
        if removing:
            to_remove.append(child)
    for el in to_remove:
        delete_element(el)


def build_appendix_a(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    title = doc.add_heading("Appendix A: File inventory", level=1)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)

    add_body(
        doc,
        "This appendix lists all repository artefacts referenced in the technical report "
        "and reproduced in Final_Report.pdf. Paths are relative to the project root.",
    )

    add_section_heading(doc, "A.0  Repository root (documentation and optional prose)")
    add_table_block(
        doc,
        "Table A.0  Repository root files",
        ["File", "Role"],
        [
            ["README.md", "Entry point: tree, run instructions, output locations"],
            ["REPORT.md", "Technical report (Markdown source)"],
            ["Final_Report.pdf", "PDF export of REPORT.md"],
            ["scripts/build_final_report_pdf.py", "Rebuild Final_Report.pdf from REPORT.md"],
            ["scripts/format_docx_appendix.py", "Rebuild this appendix in the Word report"],
            ["requirements.txt", "Python dependencies (numpy, pandas, pulp, …)"],
            ["LICENSE", "MIT licence"],
            [
                "ROBUST OPTIMISATION OF INFRASTRUCTURE SCHEDULING.docx",
                "Editable thesis / project report (this document)",
            ],
        ],
        [2.4, 3.6],
    )

    add_section_heading(doc, "A.1  Code")
    add_table_block(
        doc,
        "Table A.1  Executable models",
        ["File", "Lines (approx.)", "Function"],
        [
            [
                "Model and solutions/optimisation_model.py",
                "720",
                "Main MIP + Monte Carlo + sensitivity plots 01–05",
            ],
            [
                "O:P with pareto fronts/tradeoff_study.py",
                "650",
                "Pareto fronts, tail risk, λ trade-off plots 06–09",
            ],
        ],
        [2.8, 1.0, 2.7],
    )

    add_section_heading(doc, "A.2  Data (Baseline reference datasets/)")
    add_body(
        doc,
        "Infrastructure_Project_Dataset.xlsx is the consolidated case-study workbook. "
        "All Python scripts read the cleaned export under "
        "Baseline reference datasets/Dependency2/:",
        italic=True,
    )
    add_table_block(
        doc,
        "Table A.2  Dependency2 input files (model-ready)",
        ["File", "Records (approx.)", "Role"],
        [
            ["tasks_clean.csv", "24", "Tasks: duration, resources, road impact, DM per block"],
            ["dependencies_clean.csv", "33", "Precedence edges (FS / SS)"],
            ["cost_parameters_resources.csv", "9", "Labour unit rates, overtime, idle fractions"],
            ["cost_parameters_events.csv", "10", "Event penalties and disruption rates"],
            ["equipment_pool.csv", "5", "Shared equipment pool limits"],
            ["resource_limits.csv", "72", "Weekly crew caps"],
            ["resource_usage_daily.csv", "—", "Daily resource draw (supporting)"],
            ["resources_clean.csv", "—", "Resource master list"],
            ["risk_clean.csv", "22", "Risk events for attribution plots"],
            ["traffic_schedule_multipliers.csv", "388", "Day × block dynamic DM and delay index"],
            ["daily_schedule_grid.csv", "97", "Fine schedule grid (reference)"],
            ["time_blocks.csv", "4", "Block definitions (B1–B4)"],
            ["uncertainty_distributions.json", "—", "Monte Carlo factor specification"],
            ["cost_function_spec.json", "—", "Machine-readable cost decomposition"],
        ],
        [2.5, 1.1, 3.0],
    )

    add_section_heading(doc, "A.3  Results CSV (committed in-repo)")
    add_table_block(
        doc,
        "Table A.3  Committed result tables",
        ["File", "Location", "Description"],
        [
            ["results_summary.csv", "Model and solutions/", "Aggregate KPIs × 3 schedule variants"],
            ["actionable_decisions.csv", "Model and solutions/", "Per-task block and start deltas"],
            ["sensitivity_lambda.csv", "Model and solutions/", "λ (risk aversion) sweep"],
            ["sensitivity_gamma.csv", "Model and solutions/", "γ (traffic weight) sweep"],
            ["sensitivity_alpha.csv", "Model and solutions/", "α (cost weight) sweep"],
            ["pareto_cost_duration.csv", "O:P with pareto fronts/", "Cost–duration Pareto sweep"],
            ["pareto_cost_traffic.csv", "O:P with pareto fronts/", "Cost–traffic Pareto sweep"],
            ["risk_cost_curve.csv", "O:P with pareto fronts/", "λ vs cost percentiles"],
            ["tail_risk_summary.csv", "O:P with pareto fronts/", "P50–P99 tail metrics"],
        ],
        [2.2, 1.8, 2.5],
    )

    add_section_heading(doc, "A.4  Figures")
    add_table_block(
        doc,
        "Table A.4  Generated figures (after running both scripts)",
        ["ID", "Filename", "Topic"],
        [
            ["01", "01_distributions.png", "Monte Carlo distributions and KPI bars"],
            ["02", "02_sensitivity.png", "α, γ, λ sensitivity sweeps"],
            ["03", "03_gantt.png", "Baseline vs optimised schedule"],
            ["04", "04_actionable.png", "Task-level block and start changes"],
            ["05", "05_cascade.png", "Cascade regimes and risk categories"],
            ["06", "06_tradeoff_curves.png", "Pareto and risk–cost panels"],
            ["07", "07_stochastic_study.png", "Stochastic comparison (deep dive)"],
            ["08", "08_tail_risk.png", "Tail percentiles and CVaR"],
            ["09", "09_lambda_tradeoff.png", "Robustness premium vs λ"],
        ],
        [0.45, 2.0, 3.0],
    )

    add_body(
        doc,
        "Figures 01–05 are written to Model and solutions/plots/; figures 06–09 to "
        "O:P with pareto fronts/plots_tradeoff/. These folders are gitignored—run both "
        "Python scripts after clone to regenerate PNGs. CSV outputs remain next to each script "
        "and are tracked in git.",
    )
    add_body(
        doc,
        "Report cross-reference: Figure 1 = 01_distributions.png; Figure 2 = 03_gantt.png; "
        "Figure 3 = 06_tradeoff_curves.png; Figure 4 = 08_tail_risk.png; "
        "Figure 5 = 09_lambda_tradeoff.png. Final_Report.pdf embeds these five plus a "
        "thumbnail montage of 01–09 when plot folders are present.",
    )


def format_appendix_b_outputs(doc: Document) -> None:
    """Convert Appendix B.4 bullet list into a formatted table (idempotent)."""
    for t in doc.tables:
        if t.rows[0].cells[0].text.strip().startswith("File / script"):
            return  # already formatted

    start_p = None
    list_ps = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "B.4 Output Files":
            start_p = p
            continue
        if start_p is not None and p.style.name == "List Paragraph" and (
            ".csv" in t or ".py" in t
        ):
            list_ps.append(p)
        elif start_p is not None and list_ps and t and p.style.name != "List Paragraph":
            break

    if not list_ps:
        return

    rows = []
    for p in list_ps:
        text = p.text.strip()
        if " -" in text:
            file_part, desc = text.split(" -", 1)
        elif " —" in text:
            file_part, desc = text.split(" —", 1)
        else:
            parts = text.replace("  ", " ").split(" ", 1)
            file_part = parts[0]
            desc = parts[1] if len(parts) > 1 else ""
        rows.append([file_part.strip(), desc.strip().lstrip("-").strip()])

    for p in list_ps:
        delete_paragraph(p)

    # Insert table after B.4 heading
    idx = None
    body = doc.element.body
    for i, child in enumerate(body):
        if child.tag.endswith("p") and para_text(child).strip() == "B.4 Output Files":
            idx = i
            break
    if idx is None:
        return

    add_caption_after_element(doc, body[idx], "Table B.1  Key output artefacts")
    tbl = doc.add_table(rows=1 + len(rows), cols=2)
    format_data_table(tbl, ["File / script", "Description"], rows, [2.4, 3.6])

    # Move table to correct position (after caption, before next section)
    tbl_el = doc.tables[-1]._tbl
    body.remove(tbl_el)
    body.insert(idx + 2, tbl_el)


def add_caption_after_element(doc: Document, after_element, caption: str) -> None:
    p = OxmlElement("w:p")
    after_element.addnext(p)
    from docx.text.paragraph import Paragraph

    para = Paragraph(p, doc.element.body)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def delete_paragraph(paragraph) -> None:
    delete_element(paragraph._element)


def main() -> None:
    doc = Document(str(DOCX_PATH))
    remove_appendix_a(doc)
    build_appendix_a(doc)
    format_appendix_b_outputs(doc)
    doc.save(str(DOCX_PATH))
    print(f"Formatted appendices in {DOCX_PATH}")


if __name__ == "__main__":
    main()
